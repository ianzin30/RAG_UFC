import csv
import re
import unicodedata
from io import BytesIO, StringIO

from openpyxl import load_workbook


MAX_SPREADSHEET_ROWS = 1000
MAX_SPREADSHEET_COLUMNS = 40
MAX_SAMPLE_VALUES = 5
MAX_CELL_CHARS = 200
CSV_SNIFF_DELIMITERS = ",;\t|"
PERSON_BLOCKLIST = {
    "nome",
    "tbd",
    "total",
    "totais",
    "subtotal",
    "sub total",
    "target",
    "rpa",
    "profissionais",
    "bolsas",
    "custo dell",
    "custo total de rh mensal",
    "encargos clt",
    "descricoes",
    "descricao",
    "historico de margem para prestacao de contas dell",
    "rh direto bolsas",
    "rh direto celetistas",
    "rh indireto administrativo",
}
PERSON_PREFIX_BLOCKLIST = (
    "rh ",
    "sub total",
    "subtotal",
    "total ",
    "totais ",
    "custo ",
    "historico ",
    "margem ",
)
ROLE_BLOCKLIST = {
    "ufc",
    "s vinculo",
    "column",
    "custo dell",
    "total",
    "totais",
}


def extract_spreadsheet_markdown(file_bytes: bytes, suffix: str, file_name: str) -> str:
    normalized_suffix = (suffix or "").lower()
    if normalized_suffix == ".xlsx":
        sheets = _load_xlsx_sheets(file_bytes)
    elif normalized_suffix == ".csv":
        sheets = [_load_csv_sheet(file_bytes)]
    else:
        raise ValueError(f"Unsupported spreadsheet suffix: {suffix}")

    sheet_names = ", ".join(sheet["name"] for sheet in sheets) or "Sem abas detectadas"
    lines = [
        "Document type: spreadsheet",
        f"Spreadsheet format: {normalized_suffix.lstrip('.') or 'unknown'}",
        f"Spreadsheet file: {file_name}",
        f"Total sheets: {len(sheets)}",
        f"Sheet names: {sheet_names}",
        "",
    ]

    for sheet in sheets:
        lines.extend(_render_sheet(sheet, file_name))

    return "\n".join(line for line in lines if line is not None).strip()


def _load_xlsx_sheets(file_bytes: bytes) -> list[dict]:
    workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            normalized_row = [_normalize_cell(value) for value in row]
            if any(normalized_row):
                rows.append(normalized_row)
        sheets.append(_build_sheet(worksheet.title or "Sheet", rows))
    return sheets or [{"name": "Sheet", "headers": [], "rows": [], "omitted_rows": 0, "people_entries": []}]


def _load_csv_sheet(file_bytes: bytes) -> dict:
    text = _decode_csv_bytes(file_bytes)
    dialect = _detect_csv_dialect(text)
    reader = csv.reader(StringIO(text), dialect)
    rows = []
    for row in reader:
        normalized_row = [_normalize_cell(value) for value in row]
        if any(normalized_row):
            rows.append(normalized_row)
    return _build_sheet("CSV", rows)


def _decode_csv_bytes(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _detect_csv_dialect(text: str):
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_SNIFF_DELIMITERS)
    except csv.Error:
        return csv.excel


def _build_sheet(sheet_name: str, raw_rows: list[list[str]]) -> dict:
    if not raw_rows:
        return {"name": sheet_name, "headers": [], "rows": [], "omitted_rows": 0, "people_entries": []}

    header_row = raw_rows[0][:MAX_SPREADSHEET_COLUMNS]
    headers = _make_headers(header_row)
    data_rows = raw_rows[1 : MAX_SPREADSHEET_ROWS + 1]
    omitted_rows = max(len(raw_rows) - 1 - len(data_rows), 0)

    structured_rows = []
    for row_index, row in enumerate(data_rows, start=2):
        values = row[: len(headers)]
        if len(values) < len(headers):
            values = values + [""] * (len(headers) - len(values))
        pairs = [(header, value) for header, value in zip(headers, values) if value]
        if pairs:
            structured_rows.append({"row_number": row_index, "pairs": pairs})

    return {
        "name": sheet_name,
        "headers": headers,
        "rows": structured_rows,
        "omitted_rows": omitted_rows,
        "people_entries": _build_people_entries(sheet_name, structured_rows),
    }


def _make_headers(raw_headers: list[str]) -> list[str]:
    headers = []
    seen = {}
    for index, raw_header in enumerate(raw_headers, start=1):
        base = raw_header or f"Column_{index}"
        normalized = re.sub(r"\s+", " ", str(base)).strip() or f"Column_{index}"
        count = seen.get(normalized, 0)
        seen[normalized] = count + 1
        headers.append(normalized if count == 0 else f"{normalized}_{count + 1}")
    return headers


def _render_sheet(sheet: dict, file_name: str) -> list[str]:
    headers = sheet["headers"]
    rows = sheet["rows"]
    lines = [
        f"## Sheet: {sheet['name']}",
        "",
        f"Spreadsheet file: {file_name}",
        f"Sheet name: {sheet['name']}",
        f"Header columns: {', '.join(headers) if headers else 'No headers detected'}",
        f"Data rows indexed: {len(rows)}",
    ]

    if sheet["omitted_rows"]:
        lines.append(f"Omitted rows: {sheet['omitted_rows']}")

    column_profiles = _build_column_profiles(headers, rows)
    if column_profiles:
        lines.extend(["", "Column profiles:"])
        for profile in column_profiles:
            lines.append(f"- {profile}")

    if sheet["people_entries"]:
        lines.extend(["", "People index:"])
        for entry in sheet["people_entries"]:
            role_suffix = f" | Role: {entry['role']}" if entry.get("role") else ""
            lines.append(
                f"- Person: {entry['name']}{role_suffix} | Sheet: {sheet['name']} | Row: {entry['row_number']}"
            )

    if rows:
        lines.extend(["", "Row records:"])
        for row in rows:
            pairs_text = " | ".join(f"{header}={value}" for header, value in row["pairs"])
            lines.append(
                f"- Spreadsheet file {file_name} | Sheet {sheet['name']} | Row {row['row_number']}: {pairs_text}"
            )
    else:
        lines.extend(["", "No non-empty data rows were indexed."])

    lines.append("")
    return lines


def _build_column_profiles(headers: list[str], rows: list[dict]) -> list[str]:
    profiles = []
    for header in headers:
        values = []
        for row in rows:
            for row_header, row_value in row["pairs"]:
                if row_header == header and row_value:
                    values.append(row_value)
                    break

        if not values:
            continue

        numeric_values = [_to_number(value) for value in values]
        numeric_values = [value for value in numeric_values if value is not None]

        if len(numeric_values) >= 2:
            profiles.append(
                (
                    f"{header}: numeric count={len(numeric_values)}, "
                    f"sum={_format_number(sum(numeric_values))}, "
                    f"min={_format_number(min(numeric_values))}, "
                    f"max={_format_number(max(numeric_values))}"
                )
            )
            continue

        samples = []
        seen = set()
        for value in values:
            sample = value[:80]
            if sample in seen:
                continue
            seen.add(sample)
            samples.append(sample)
            if len(samples) >= MAX_SAMPLE_VALUES:
                break

        if samples:
            profiles.append(f"{header}: sample values={'; '.join(samples)}")

    return profiles


def _build_people_entries(sheet_name: str, rows: list[dict]) -> list[dict]:
    entries = []
    seen = set()
    for row in rows:
        entry = _extract_person_entry(sheet_name, row)
        if not entry:
            continue
        key = (entry["name"], entry["role"], entry["row_number"], entry["sheet_name"])
        if key in seen:
            continue
        seen.add(key)
        entries.append(entry)
    return entries


def _extract_person_entry(sheet_name: str, row: dict) -> dict | None:
    if not row.get("pairs"):
        return None

    first_header, first_value = row["pairs"][0]
    if not _looks_like_person_name(first_value):
        return None

    role = _extract_role_from_pairs(row["pairs"][1:])
    return {
        "name": first_value,
        "role": role,
        "row_number": row["row_number"],
        "sheet_name": sheet_name,
        "source_header": first_header,
    }


def _extract_role_from_pairs(pairs: list[tuple[str, str]]) -> str | None:
    for header, value in pairs:
        if _looks_like_role(header, value):
            return value
    return None


def _looks_like_person_name(value: str) -> bool:
    if not value:
        return False

    text = re.sub(r"\s+", " ", value).strip()
    if len(text) < 5 or any(char.isdigit() for char in text):
        return False

    if not any(char.islower() for char in text):
        return False

    normalized = _normalize_identifier(text)
    if not normalized or normalized in PERSON_BLOCKLIST:
        return False
    if any(normalized.startswith(prefix) for prefix in PERSON_PREFIX_BLOCKLIST):
        return False

    tokens = re.findall(r"[A-Za-zÀ-ÿ]+", text)
    if len(tokens) < 2:
        return False

    return True


def _looks_like_role(header: str, value: str) -> bool:
    if not value:
        return False

    text = re.sub(r"\s+", " ", value).strip()
    normalized = _normalize_identifier(text)
    if not normalized or normalized in ROLE_BLOCKLIST:
        return False

    if _to_number(text) is not None:
        return False

    if re.search(r"\d{4}-\d{2}-\d{2}", text):
        return False

    tokens = re.findall(r"[A-Za-zÀ-ÿ]+", text)
    if not tokens or len(tokens) > 12:
        return False

    normalized_header = _normalize_identifier(header)
    if normalized_header.startswith("column_") and len(tokens) <= 1:
        return False

    return True


def _normalize_identifier(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", (text or "").lower())
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def _normalize_cell(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()

    text = re.sub(r"\s+", " ", str(value)).strip()
    if len(text) > MAX_CELL_CHARS:
        return text[: MAX_CELL_CHARS - 3] + "..."
    return text


def _to_number(value: str) -> float | None:
    if not value:
        return None

    cleaned = value.strip().replace("R$", "").replace("%", "").replace("\u00a0", " ")
    cleaned = cleaned.replace(" ", "")
    if not cleaned:
        return None

    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")

    try:
        return float(cleaned)
    except ValueError:
        return None


def _format_number(value: float) -> str:
    if value.is_integer():
        return str(int(value))
    return f"{value:.2f}"
