import sys
from io import BytesIO
from pathlib import Path
from textwrap import dedent

from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnableLambda
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from application.service.collections.contracts import NormalizedDocument, QueryPlan, SpreadsheetModel
from application.service.collections.document_normalizer import DocumentNormalizer
from application.service.collections.repository import CollectionRepository
from application.service.bot_support.extraction_selection import (
    EXTRACTION_METHOD_DOC_CSV,
    build_extraction_choice_prompt,
    format_extraction_method,
    normalize_extraction_method,
)
from application.service.google_drive import GoogleDriveService
from application.service.rag import RAGService
from application.service.retrieval.coordinator import RetrievalCoordinator
from application.service.retrieval.query_planner import QueryPlanner
from application.service.spreadsheets.markdown import (
    extract_spreadsheet_csv_children,
    extract_spreadsheet_markdown,
    extract_spreadsheet_markdown_via_csv,
)
from application.service.spreadsheets.normalizer import SpreadsheetNormalizer


def test_extract_spreadsheet_markdown_emits_people_index():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RH"
    sheet.append(["RH DIRETO - BOLSAS", "Column_2", "Column_3"])
    sheet.append(["NOME", "Cargo/Função", "Início"])
    sheet.append(["Fernando Antonio Mota Trinta", "Consultor em Arquitetura de Software", "1.0"])
    sheet.append(["Maria Letícia dos Santos Dantas", "Testador 2", "1.0"])
    sheet.append(["KARLA CRISTINA SOARES SOUZA", "Gerente de projetos", "1.0"])
    sheet.append(["RH INDIRETO - ADMINISTRATIVO", "", ""])
    sheet.append(["TOTAL", "", ""])

    planning_sheet = workbook.create_sheet("PLANEJADO")
    planning_sheet.append(["INSTITUTO/IES: UFC", "Column_2", "Column_3"])
    planning_sheet.append(["Serviços outros", "", ""])
    planning_sheet.append(["Licenças Gsuit", "", ""])
    planning_sheet.append(["Fundo de reserva", "Demais custos de P&D", ""])

    buffer = BytesIO()
    workbook.save(buffer)

    markdown = extract_spreadsheet_markdown(buffer.getvalue(), ".xlsx", "PlanilhaFinanceira.xlsx")

    assert "People index:" in markdown
    assert "Person: Fernando Antonio Mota Trinta | Role: Consultor em Arquitetura de Software | Sheet: RH | Row: 3" in markdown
    assert "Person: Maria Letícia dos Santos Dantas | Role: Testador 2 | Sheet: RH | Row: 4" in markdown
    assert "Person: KARLA CRISTINA SOARES SOUZA | Role: Gerente de projetos | Sheet: RH | Row: 5" in markdown
    assert "Person: RH INDIRETO - ADMINISTRATIVO" not in markdown
    assert "Person: TOTAL" not in markdown
    assert "Person: Serviços outros" not in markdown


def test_extract_spreadsheet_markdown_via_csv_preserves_sheet_names_and_people_index():
    workbook = Workbook()
    rh_sheet = workbook.active
    rh_sheet.title = "RH"
    rh_sheet.append(["NOME", "Cargo/Função", "Início"])
    rh_sheet.append(["Fernando Antonio Mota Trinta", "Consultor em Arquitetura de Software", "1.0"])
    costs_sheet = workbook.create_sheet("CUSTOS")
    costs_sheet.append(["Categoria", "Valor"])
    costs_sheet.append(["Serviços Técnicos", "2500"])

    buffer = BytesIO()
    workbook.save(buffer)

    markdown = extract_spreadsheet_markdown_via_csv(buffer.getvalue(), ".xlsx", "PlanilhaFinanceira.xlsx")

    assert "Spreadsheet route: csv-per-sheet" in markdown
    assert "Sheet names: RH, CUSTOS" in markdown
    assert "## Sheet: RH" in markdown
    assert "## Sheet: CUSTOS" in markdown
    assert "Person: Fernando Antonio Mota Trinta | Role: Consultor em Arquitetura de Software | Sheet: RH | Row: 2" in markdown


def test_extract_spreadsheet_csv_children_returns_one_csv_text_per_sheet():
    workbook = Workbook()
    rh_sheet = workbook.active
    rh_sheet.title = "RH"
    rh_sheet.append(["NOME", "Cargo/Função"])
    rh_sheet.append(["Fernando Antonio Mota Trinta", "Consultor em Arquitetura de Software"])
    planned_sheet = workbook.create_sheet("PLANEJADO")
    planned_sheet.append(["Categoria", "Valor"])
    planned_sheet.append(["Equipamentos", "31300"])

    buffer = BytesIO()
    workbook.save(buffer)

    children = extract_spreadsheet_csv_children(buffer.getvalue(), ".xlsx")

    assert [child["sheet_name"] for child in children] == ["RH", "PLANEJADO"]
    assert "NOME,Cargo/Função" in children[0]["csv_text"]
    assert "Categoria,Valor" in children[1]["csv_text"]


def test_spreadsheet_normalizer_parses_structured_markdown_and_preserves_people_rows():
    normalizer = SpreadsheetNormalizer()
    model = normalizer.normalize("PlanilhaFinanceira.xlsx", _legacy_spreadsheet_markdown(), "xlsx")

    assert model.sheets
    rh_sheet = model.sheets[0]
    row_entities = [row.entity.name for row in rh_sheet.rows if row.entity]

    assert row_entities == [
        "Fernando Antonio Mota Trinta",
        "Maria Letícia dos Santos Dantas",
        "Paulo Antonio Rego",
    ]
    assert [person.name for person in rh_sheet.people] == row_entities


def test_spreadsheet_normalizer_parses_docling_pure_markdown_tables():
    normalizer = SpreadsheetNormalizer()
    model = normalizer.normalize("PlanilhaFinanceira.xlsx", _docling_pure_table_markdown(), "xlsx")

    assert model.sheets
    sheet = model.sheets[0]
    assert [person.name for person in sheet.people][:2] == [
        "Fernando Antonio Mota Trinta",
        "Maria Letícia dos Santos Dantas",
    ]


def test_doc_csv_selection_and_google_drive_method_normalization_accept_doc_csv():
    service = GoogleDriveService()

    assert normalize_extraction_method("doc-csv") == EXTRACTION_METHOD_DOC_CSV
    assert normalize_extraction_method("doc csv") == EXTRACTION_METHOD_DOC_CSV
    assert normalize_extraction_method("doc_csv") == EXTRACTION_METHOD_DOC_CSV
    assert normalize_extraction_method("4") == EXTRACTION_METHOD_DOC_CSV
    assert service._normalize_extraction_method("doc-csv") == EXTRACTION_METHOD_DOC_CSV
    assert format_extraction_method(EXTRACTION_METHOD_DOC_CSV) == "Doc CSV"
    assert "`doc-csv`" in build_extraction_choice_prompt()


def test_doc_csv_spreadsheet_document_persists_route_in_summary_and_sidecar(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RH"
    sheet.append(["NOME", "Cargo/Função"])
    sheet.append(["Paulo Antonio Rego", "Coordenador do Projeto"])

    buffer = BytesIO()
    workbook.save(buffer)

    markdown = extract_spreadsheet_markdown_via_csv(buffer.getvalue(), ".xlsx", "PlanilhaFinanceira.xlsx")
    normalizer = SpreadsheetNormalizer()
    model = normalizer.normalize("PlanilhaFinanceira.xlsx", markdown, "xlsx")
    document = DocumentNormalizer().normalize_drive_document(
        index=1,
        drive_file={"id": "spreadsheet-1", "name": "PlanilhaFinanceira.xlsx", "mimeType": "application/vnd.ms-excel"},
        extraction_method=EXTRACTION_METHOD_DOC_CSV,
        content_markdown=markdown,
        spreadsheet_model=model,
    )

    repository = CollectionRepository(tmp_path)
    repository.save_collection("doc_csv_contract", "google_drive", EXTRACTION_METHOD_DOC_CSV, [document])
    loaded = repository.load_collection("doc_csv_contract")

    loaded_document = loaded.documents[0]
    assert loaded.manifest.extraction_method == EXTRACTION_METHOD_DOC_CSV
    assert loaded_document.extraction_method == EXTRACTION_METHOD_DOC_CSV
    assert "Spreadsheet route: csv-per-sheet" in loaded_document.summary
    assert "Spreadsheet route: csv-per-sheet" in loaded_document.content_text
    assert "Spreadsheet route: csv-per-sheet" in loaded_document.structured_data["summary_lines"]


def test_doc_csv_xlsx_builds_visible_parent_and_internal_sheet_children():
    workbook = Workbook()
    rh_sheet = workbook.active
    rh_sheet.title = "RH"
    rh_sheet.append(["NOME", "Cargo/Função"])
    rh_sheet.append(["Paulo Antonio Rego", "Coordenador do Projeto"])
    costs_sheet = workbook.create_sheet("CUSTOS")
    costs_sheet.append(["Categoria", "Valor"])
    costs_sheet.append(["Equipamentos", "31300"])

    buffer = BytesIO()
    workbook.save(buffer)

    service = GoogleDriveService()
    documents = service._build_doc_csv_spreadsheet_documents(
        index=6,
        drive_file={"id": "drive-sheet-1", "name": "PlanilhaFinanceira.xlsx", "mimeType": "application/vnd.google-apps.spreadsheet"},
        extraction_method=EXTRACTION_METHOD_DOC_CSV,
        file_bytes=buffer.getvalue(),
        suffix=".xlsx",
    )

    parent = documents[0]
    children = documents[1:]

    assert parent.display_name == "PlanilhaFinanceira.xlsx"
    assert parent.catalog_visibility == "visible"
    assert parent.component_kind == "spreadsheet_parent"
    assert len(children) == 2
    assert all(child.catalog_visibility == "internal" for child in children)
    assert all(child.parent_document_id == parent.document_id for child in children)
    assert all(child.logical_item_id == parent.document_id for child in children)
    assert [child.component_name for child in children] == ["RH", "CUSTOS"]
    assert children[0].display_name == "PlanilhaFinanceira.xlsx::RH.csv"
    assert children[0].document_type == "text_document"
    assert "Workbook: PlanilhaFinanceira.xlsx" in children[0].content_text
    assert "Sheet: RH" in children[0].content_text
    assert "```csv" in children[0].content_markdown


def test_doc_csv_direct_csv_markdown_marks_direct_route():
    csv_bytes = (
        "NOME,Cargo/Função,Início\n"
        "Fernando Antonio Mota Trinta,Consultor em Arquitetura de Software,1.0\n"
    ).encode("utf-8")

    markdown = extract_spreadsheet_markdown(
        csv_bytes,
        ".csv",
        "PlanilhaFinanceira.csv",
        route_label="csv-direct",
    )

    assert "Spreadsheet route: csv-direct" in markdown
    assert "## Sheet: CSV" in markdown


def test_collection_repository_loads_manifest_backed_collection_without_markdown_inference(tmp_path):
    repository = CollectionRepository(tmp_path)
    document = NormalizedDocument(
        document_id="01_contract",
        display_name="Canonical.pdf",
        document_type="pdf",
        source_kind="google_drive",
        source_metadata={"id": "123"},
        extraction_method="docling-puro",
        content_markdown_path="",
        content_text="Texto canônico do documento.",
        summary="Texto canônico do documento.",
        structured_data=None,
        content_markdown="# WrongHeading\n\nThis markdown header must not become the document identity.",
    )

    repository.save_collection("contract_test", "google_drive", "docling-puro", [document])
    loaded = repository.load_collection("contract_test")

    assert loaded.manifest.collection_name == "contract_test"
    assert loaded.documents[0].display_name == "Canonical.pdf"
    assert loaded.documents[0].document_type == "pdf"


def test_collection_repository_rejects_legacy_markdown_only_collection(tmp_path):
    repository = CollectionRepository(tmp_path)
    legacy_path = tmp_path / "data" / "collections" / "legacy"
    legacy_path.mkdir(parents=True, exist_ok=True)
    (legacy_path / "01_legacy.md").write_text("# Legacy", encoding="utf-8")

    try:
        repository.load_collection("legacy")
    except Exception as exc:
        assert "manifest.json" in str(exc)
    else:
        raise AssertionError("Legacy collections must be rejected without manifest.json")


def test_query_planner_routes_inventory_multi_doc_and_spreadsheet_entity_questions():
    planner = QueryPlanner()
    catalog = [
        planner.build_catalog_entry("01_afo_pdf", "AFO.pdf", "pdf", ""),
        planner.build_catalog_entry("02_afo_xlsx", "AFO.xlsx", "spreadsheet", ""),
        planner.build_catalog_entry("03_planilha", "PlanilhaFinanceira.xlsx", "spreadsheet", ""),
    ]

    inventory = planner.plan(
        "quantos arquivos voce processou?",
        "quantos arquivos voce processou?",
        catalog,
        intent_decision={"intent": "inventory_query", "confidence": 0.93},
    )
    comparison = planner.plan(
        "compare AFO.pdf e AFO.xlsx",
        "compare AFO.pdf e AFO.xlsx",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.91},
    )
    spreadsheet = planner.plan(
        "quais pessoas aparecem na PlanilhaFinanceira.xlsx?",
        "quais pessoas aparecem na PlanilhaFinanceira.xlsx?",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.95},
    )

    assert inventory.intent == "inventory"
    assert inventory.inventory_mode == "count"
    assert comparison.intent == "multi_document"
    assert comparison.target_document_ids == ["01_afo_pdf", "02_afo_xlsx"]
    assert spreadsheet.intent == "spreadsheet_entity"
    assert spreadsheet.target_document_ids == ["03_planilha"]


def test_query_planner_sets_doc_csv_hybrid_retrieval_profile_by_question_shape():
    planner = QueryPlanner()
    catalog = [
        planner.build_catalog_entry(
            "06_planilha",
            "PlanilhaFinanceira.xlsx",
            "spreadsheet",
            "Resumo da planilha.",
            backing_document_ids=["06_planilha", "06_planilha__RH", "06_planilha__CUSTOS"],
            primary_document_id="06_planilha",
            child_document_ids=["06_planilha__RH", "06_planilha__CUSTOS"],
            supports_doc_csv_hybrid=True,
            child_components=[
                {"document_id": "06_planilha__RH", "component_name": "RH"},
                {"document_id": "06_planilha__CUSTOS", "component_name": "CUSTOS"},
            ],
            contained_document_types=["spreadsheet", "text_document"],
            match_names=["PlanilhaFinanceira.xlsx::RH.csv", "PlanilhaFinanceira.xlsx::CUSTOS.csv"],
        )
    ]

    broad = planner.plan(
        "me fale sobre o arquivo PlanilhaFinanceira.xlsx",
        "me fale sobre o arquivo PlanilhaFinanceira.xlsx",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.93},
    )
    people = planner.plan(
        "quais pessoas aparecem na PlanilhaFinanceira.xlsx?",
        "quais pessoas aparecem na PlanilhaFinanceira.xlsx?",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.95},
    )
    sheet = planner.plan(
        "o que existe na aba RH da PlanilhaFinanceira.xlsx?",
        "o que existe na aba RH da PlanilhaFinanceira.xlsx?",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.94},
    )

    assert broad.intent == "single_document"
    assert broad.retrieval_profile == "parent_overview_first"
    assert broad.target_document_ids == ["06_planilha", "06_planilha__RH", "06_planilha__CUSTOS"]
    assert people.retrieval_profile == "csv_child_first"
    assert people.target_document_ids == ["06_planilha__RH", "06_planilha__CUSTOS", "06_planilha"]
    assert sheet.retrieval_profile == "csv_child_first"


def test_query_planner_treats_logical_folder_as_single_item_and_targets_all_backing_documents():
    planner = QueryPlanner()
    catalog = [
        planner.build_catalog_entry(
            "group_docs",
            "docs.example.com/api/",
            "text_document",
            "Pasta lógica com páginas da documentação.",
            backing_document_ids=["01_intro", "02_auth", "03_limits"],
            logical_item_kind="folder",
            member_count=3,
            match_names=["intro.md", "authentication.md", "rate-limits.md"],
            contained_document_types=["text_document"],
        ),
        planner.build_catalog_entry("04_afo", "AFO.pdf", "pdf", ""),
    ]

    plan = planner.plan(
        "me fale sobre a authentication.md",
        "me fale sobre a authentication.md",
        catalog,
        intent_decision={"intent": "general", "confidence": 0.91},
    )
    inventory = planner.plan(
        "quais arquivos estao presentes?",
        "quais arquivos estao presentes?",
        catalog,
        intent_decision={"intent": "inventory_query", "confidence": 0.94},
    )
    inventory_answer = planner.answer_inventory(inventory, catalog)

    assert plan.intent == "single_document"
    assert plan.target_document_ids == ["01_intro", "02_auth", "03_limits"]
    assert "Encontrei 2 itens carregados:" in inventory_answer
    assert "docs.example.com/api/ (pasta com 3 arquivos, documentos de texto)" in inventory_answer


def test_retrieval_coordinator_retrieves_people_from_normalized_spreadsheet_model():
    normalizer = SpreadsheetNormalizer()
    model = normalizer.normalize("PlanilhaFinanceira.xlsx", _legacy_spreadsheet_markdown(), "xlsx")
    document = _build_spreadsheet_document("01_planilha", "PlanilhaFinanceira.xlsx", model)
    coordinator = RetrievalCoordinator(
        RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=100,
            separators=["\n\n", "\n", ". ", "? ", "! ", "; ", ": ", "  ", " ", ""],
        )
    )

    chunks = coordinator._build_spreadsheet_chunks(document, SpreadsheetModel.from_dict(document.structured_data))
    docs = coordinator.retrieve(
        plan=QueryPlan(
            intent="spreadsheet_entity",
            target_document_ids=["01_planilha"],
            requested_document_type="spreadsheet",
            needs_structured_lookup=True,
            resolved_query="quais sao alguns trabalhadores mencionados na PlanilhaFinanceira.xlsx?",
        ),
        vector_store=None,
        retriever=None,
        spreadsheet_chunk_index={"01_planilha": chunks},
        document_catalog=[
            {
                "document_id": "01_planilha",
                "display_name": "PlanilhaFinanceira.xlsx",
                "document_type": "spreadsheet",
            }
        ],
    )

    names = [doc.metadata.get("entity_name") for doc in docs if doc.metadata.get("entity_name")]
    assert names[:3] == [
        "Fernando Antonio Mota Trinta",
        "Maria Letícia dos Santos Dantas",
        "Paulo Antonio Rego",
    ]


def test_doc_csv_parent_spreadsheet_chunks_are_overview_only_and_children_use_generic_text():
    normalizer = SpreadsheetNormalizer()
    model = normalizer.normalize("PlanilhaFinanceira.xlsx", _legacy_spreadsheet_markdown(), "xlsx")
    parent_document = NormalizedDocument(
        document_id="06_planilha",
        display_name="PlanilhaFinanceira.xlsx",
        document_type="spreadsheet",
        source_kind="google_drive",
        source_metadata={"id": "06_planilha"},
        extraction_method=EXTRACTION_METHOD_DOC_CSV,
        content_markdown_path="",
        content_text=model.to_text(),
        summary="Resumo da planilha",
        structured_data=model.to_dict(),
        content_markdown=_legacy_spreadsheet_markdown(),
        logical_item_id="06_planilha",
        logical_item_name="PlanilhaFinanceira.xlsx",
        logical_item_kind="file",
        catalog_visibility="visible",
        component_kind="spreadsheet_parent",
    )
    child_document = NormalizedDocument(
        document_id="06_planilha__RH",
        display_name="PlanilhaFinanceira.xlsx::RH.csv",
        document_type="text_document",
        source_kind="google_drive",
        source_metadata={"id": "06_planilha", "sheet_name": "RH"},
        extraction_method=EXTRACTION_METHOD_DOC_CSV,
        content_markdown_path="",
        content_text="Workbook: PlanilhaFinanceira.xlsx\nSheet: RH\nSpreadsheet route: csv-per-sheet\n\nNOME,Cargo/Função\nPaulo Antonio Rego,Coordenador do Projeto",
        summary="Workbook: PlanilhaFinanceira.xlsx Sheet: RH Spreadsheet route: csv-per-sheet",
        structured_data=None,
        content_markdown="```csv\nNOME,Cargo/Função\nPaulo Antonio Rego,Coordenador do Projeto\n```",
        logical_item_id="06_planilha",
        logical_item_name="PlanilhaFinanceira.xlsx",
        logical_item_kind="file",
        catalog_visibility="internal",
        parent_document_id="06_planilha",
        component_kind="csv_sheet_child",
        component_name="RH",
    )

    coordinator = RetrievalCoordinator(
        RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=100,
            separators=["\n\n", "\n", ". ", "? ", "! ", "; ", ": ", "  ", " ", ""],
        )
    )

    parent_chunks = coordinator._build_spreadsheet_chunks(
        parent_document,
        SpreadsheetModel.from_dict(parent_document.structured_data),
    )
    child_chunks = coordinator._build_generic_chunks(child_document)

    assert {chunk.metadata.get("chunk_kind") for chunk in parent_chunks} <= {"summary", "sheet_summary"}
    assert "row_record" not in {chunk.metadata.get("chunk_kind") for chunk in parent_chunks}
    assert child_chunks[0].metadata["component_kind"] == "csv_sheet_child"
    assert child_chunks[0].metadata["component_name"] == "RH"
    assert child_chunks[0].metadata["parent_document_id"] == "06_planilha"


def test_rag_build_document_catalog_hides_doc_csv_children_and_keeps_parent_visible():
    rag = _build_rag_stub()
    docs = [
        NormalizedDocument(
            document_id="06_planilha",
            display_name="PlanilhaFinanceira.xlsx",
            document_type="spreadsheet",
            source_kind="google_drive",
            source_metadata={"id": "06_planilha"},
            extraction_method=EXTRACTION_METHOD_DOC_CSV,
            content_markdown_path="",
            content_text="Resumo",
            summary="Resumo",
            structured_data={"summary_lines": ["Resumo"]},
            content_markdown="# Parent",
            logical_item_id="06_planilha",
            logical_item_name="PlanilhaFinanceira.xlsx",
            logical_item_kind="file",
            catalog_visibility="visible",
            component_kind="spreadsheet_parent",
        ),
        NormalizedDocument(
            document_id="06_planilha__RH",
            display_name="PlanilhaFinanceira.xlsx::RH.csv",
            document_type="text_document",
            source_kind="google_drive",
            source_metadata={"id": "06_planilha", "sheet_name": "RH"},
            extraction_method=EXTRACTION_METHOD_DOC_CSV,
            content_markdown_path="",
            content_text="CSV",
            summary="CSV",
            structured_data=None,
            content_markdown="```csv\n...\n```",
            logical_item_id="06_planilha",
            logical_item_name="PlanilhaFinanceira.xlsx",
            logical_item_kind="file",
            catalog_visibility="internal",
            parent_document_id="06_planilha",
            component_kind="csv_sheet_child",
            component_name="RH",
        ),
    ]

    catalog = rag._build_document_catalog(docs)
    rag.document_catalog = catalog

    assert len(catalog) == 1
    assert catalog[0]["display_name"] == "PlanilhaFinanceira.xlsx"
    assert catalog[0]["primary_document_id"] == "06_planilha"
    assert catalog[0]["child_document_ids"] == ["06_planilha__RH"]
    assert catalog[0]["supports_doc_csv_hybrid"] is True
    assert rag._format_target_document_names(["06_planilha__RH"]) == "PlanilhaFinanceira.xlsx"


def test_ask_question_answers_collection_inventory_from_catalog_without_llm():
    rag = _build_rag_stub()
    rag.retriever = object()
    rag.answer_chain = _FailingChain()
    rag.small_talk_chain = _FailingChain()
    rag.intent_decider_chain = _StaticChain('{"intent":"inventory_query","confidence":0.93}')
    rag.collection_name = "documentos de teste"
    rag.document_catalog = [
        rag.query_planner.build_catalog_entry("01_afo", "AFO.pdf", "pdf", ""),
        rag.query_planner.build_catalog_entry("02_saar_pdf", "SAAR.pdf", "pdf", ""),
        rag.query_planner.build_catalog_entry("03_rack", "RACK.pdf", "pdf", ""),
        rag.query_planner.build_catalog_entry("04_saar_xlsx", "SAAR.xlsx", "spreadsheet", ""),
        rag.query_planner.build_catalog_entry("05_planilha", "PlanilhaFinanceira.xlsx", "spreadsheet", ""),
    ]

    answer = rag.ask_question("quais seriam os documentos?")

    assert "Encontrei 5 arquivos carregados:" in answer
    assert "AFO.pdf" in answer
    assert "SAAR.xlsx" in answer
    assert "PlanilhaFinanceira.xlsx" in answer


def test_rag_build_document_catalog_groups_logical_folder_as_single_catalog_item():
    rag = _build_rag_stub()
    docs = [
        NormalizedDocument(
            document_id="01_intro",
            display_name="intro.md",
            document_type="text_document",
            source_kind="scraped_web",
            source_metadata={"url": "https://docs.example.com/api/intro"},
            extraction_method="firecrawl",
            content_markdown_path="",
            content_text="Intro",
            summary="Intro",
            structured_data=None,
            content_markdown="# Intro",
            logical_item_id="group_docs",
            logical_item_name="docs.example.com/api/",
            logical_item_kind="folder",
        ),
        NormalizedDocument(
            document_id="02_auth",
            display_name="authentication.md",
            document_type="text_document",
            source_kind="scraped_web",
            source_metadata={"url": "https://docs.example.com/api/authentication"},
            extraction_method="firecrawl",
            content_markdown_path="",
            content_text="Authentication",
            summary="Authentication",
            structured_data=None,
            content_markdown="# Authentication",
            logical_item_id="group_docs",
            logical_item_name="docs.example.com/api/",
            logical_item_kind="folder",
        ),
    ]

    catalog = rag._build_document_catalog(docs)

    assert len(catalog) == 1
    assert catalog[0]["document_id"] == "group_docs"
    assert catalog[0]["display_name"] == "docs.example.com/api/"
    assert catalog[0]["logical_item_kind"] == "folder"
    assert catalog[0]["member_count"] == 2
    assert catalog[0]["backing_document_ids"] == ["01_intro", "02_auth"]


def test_parse_intent_decision_falls_back_to_general_on_invalid_json():
    rag = _build_rag_stub()

    assert rag._parse_intent_decision("") == {"intent": "general", "confidence": 0.0}
    assert rag._parse_intent_decision("not json") == {"intent": "general", "confidence": 0.0}
    assert rag._parse_intent_decision('{"intent":"unknown","confidence":"bad"}') == {
        "intent": "general",
        "confidence": 0.0,
    }
    assert rag._parse_intent_decision('```json\n{"intent":"inventory_query","confidence":1.4}\n```') == {
        "intent": "inventory_query",
        "confidence": 1.0,
    }


def test_intent_decider_prompt_formats_without_unescaped_json_variables():
    rag = _build_rag_stub()
    rag.llm = RunnableLambda(lambda value: value)
    rag.intent_decider_llm = RunnableLambda(lambda _value: '{"intent":"inventory_query","confidence":0.93}')

    rag._build_chains()

    answer = rag.intent_decider_chain.invoke(
        {
            "chat_history": "Sem conversa anterior.",
            "question": "quais arquivos estao presentes?",
            "resolved_question": "quais arquivos estao presentes?",
        }
    )

    parsed = StrOutputParser().invoke(answer)
    assert '"intent"' in parsed


def _build_rag_stub() -> RAGService:
    rag = RAGService.__new__(RAGService)
    rag.project_root = Path("/Users/fernandotrinta/Desktop/RAG")
    rag.collections_root = rag.project_root / "data" / "collections"
    rag.query_planner = QueryPlanner()
    rag.retrieval_coordinator = None
    rag.vector_store = None
    rag.retriever = None
    rag.answer_chain = None
    rag.small_talk_chain = None
    rag.query_rewrite_chain = None
    rag.intent_decider_chain = None
    rag.loaded_collection_id = None
    rag.collection_name = None
    rag.collection_manifest = None
    rag.normalized_documents = []
    rag.document_catalog = []
    rag.spreadsheet_chunk_index = {}
    return rag


class _FailingChain:
    def invoke(self, *_args, **_kwargs):
        raise AssertionError("LLM chain should not be invoked for collection inventory questions.")


class _StaticChain:
    def __init__(self, value: str):
        self.value = value

    def invoke(self, *_args, **_kwargs):
        return self.value


def _build_spreadsheet_document(document_id: str, display_name: str, model: SpreadsheetModel) -> NormalizedDocument:
    content_text = model.to_text()
    return NormalizedDocument(
        document_id=document_id,
        display_name=display_name,
        document_type="spreadsheet",
        source_kind="google_drive",
        source_metadata={"id": document_id},
        extraction_method="docling",
        content_markdown_path="",
        content_text=content_text,
        summary=content_text.split("\n", 1)[0],
        structured_data=model.to_dict(),
        content_markdown=_legacy_spreadsheet_markdown(),
    )


def _legacy_spreadsheet_markdown() -> str:
    return dedent(
        """
        Document type: spreadsheet
        Spreadsheet format: xlsx
        Spreadsheet file: PlanilhaFinanceira.xlsx
        Total sheets: 1
        Sheet names: RH
        ## Sheet: RH
        Spreadsheet file: PlanilhaFinanceira.xlsx
        Sheet name: RH
        Header columns: RH DIRETO - BOLSAS, Column_2, Column_3
        Data rows indexed: 3
        Row records:
        - Spreadsheet file PlanilhaFinanceira.xlsx | Sheet RH | Row 7: RH DIRETO - BOLSAS=Fernando Antonio Mota Trinta | Column_2=Consultor em Arquitetura de Software | Column_3=1.0
        - Spreadsheet file PlanilhaFinanceira.xlsx | Sheet RH | Row 8: RH DIRETO - BOLSAS=Maria Letícia dos Santos Dantas | Column_2=Testador 2 | Column_3=1.0
        - Spreadsheet file PlanilhaFinanceira.xlsx | Sheet RH | Row 9: RH DIRETO - BOLSAS=Paulo Antonio Rego | Column_2=Coordenador do Projeto | Column_3=1.0
        """
    ).strip()


def _docling_pure_table_markdown() -> str:
    return dedent(
        """
        | RH DIRETO - BOLSAS              | RH DIRETO - BOLSAS                   | RH DIRETO - BOLSAS |
        |---------------------------------|--------------------------------------|--------------------|
        | NOME                            | Cargo/Função                         | Início             |
        | Fernando Antonio Mota Trinta    | Consultor em Arquitetura de Software | 1.0                |
        | Maria Letícia dos Santos Dantas | Testador 2                           | 1.0                |
        | Paulo Antonio Rego              | Coordenador do Projeto               | 1.0                |
        """
    ).strip()
