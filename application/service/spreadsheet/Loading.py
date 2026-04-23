"""Workbook and CSV loading helpers for spreadsheet markdown extraction."""
# Simple: Read Excel and CSV files

from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

from .Constants import CSV_SNIFF_DELIMITERS, MAX_SPREADSHEET_COLUMNS, MAX_SPREADSHEET_ROWS
from .Models import RowRecord, SheetData
from .Normalization import make_headers, normalize_cell
from .People import build_people_entries


# Esta porta decide como carregar o arquivo conforme a extensao recebida.
def load_spreadsheet_sheets(file_bytes: bytes, suffix: str) -> list[SheetData]:
    normalized_suffix = (suffix or "").lower()
    if normalized_suffix == ".xlsx":
        return load_xlsx_sheets(file_bytes)
    if normalized_suffix == ".csv":
        return [load_csv_sheet(file_bytes)]
    raise ValueError(f"Unsupported spreadsheet suffix: {suffix}")


# Este caminho le cada aba do Excel e normaliza as linhas nao vazias.
def load_xlsx_sheets(file_bytes: bytes) -> list[SheetData]:
    workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            normalized_row = [normalize_cell(value) for value in row]
            if any(normalized_row):
                rows.append(normalized_row)
        sheets.append(build_sheet(worksheet.title or "Sheet", rows))
    return sheets or [empty_sheet("Sheet")]


# Este caminho trata CSV como uma planilha de aba unica.
def load_csv_sheet(file_bytes: bytes) -> SheetData:
    text = decode_csv_bytes(file_bytes)
    dialect = detect_csv_dialect(text)
    reader = csv.reader(StringIO(text), dialect)
    rows = []
    for row in reader:
        normalized_row = [normalize_cell(value) for value in row]
        if any(normalized_row):
            rows.append(normalized_row)
    return build_sheet("CSV", rows)


# Esta decodificacao tenta as codificacoes mais comuns antes de cair no modo tolerante.
def decode_csv_bytes(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


# Este detector tenta descobrir delimitador automaticamente sem exigir configuracao manual.
def detect_csv_dialect(text: str):
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_SNIFF_DELIMITERS)
    except csv.Error:
        return csv.excel


# Esta montagem transforma linhas cruas em estrutura pronta para renderizacao e busca.
def build_sheet(sheet_name: str, raw_rows: list[list[str]]) -> SheetData:
    if not raw_rows:
        return empty_sheet(sheet_name)

    header_row = raw_rows[0][:MAX_SPREADSHEET_COLUMNS]
    headers = make_headers(header_row)
    data_rows = raw_rows[1 : MAX_SPREADSHEET_ROWS + 1]
    omitted_rows = max(len(raw_rows) - 1 - len(data_rows), 0)

    structured_rows: list[RowRecord] = []
    for row_index, row in enumerate(data_rows, start=2):
        values = row[: len(headers)]
        if len(values) < len(headers):
            values = values + [""] * (len(headers) - len(values))
        pairs = [(header, value) for header, value in zip(headers, values) if value]
        if pairs:
            structured_rows.append(RowRecord(row_number=row_index, pairs=pairs))

    return SheetData(
        name=sheet_name,
        headers=headers,
        rows=structured_rows,
        omitted_rows=omitted_rows,
        people_entries=build_people_entries(sheet_name, structured_rows),
    )


# Esta estrutura vazia evita tratar planilhas sem dados como caso especial em todo o fluxo.
def empty_sheet(sheet_name: str) -> SheetData:
    return SheetData(
        name=sheet_name,
        headers=[],
        rows=[],
        omitted_rows=0,
        people_entries=[],
    )
